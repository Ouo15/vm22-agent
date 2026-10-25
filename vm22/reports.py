"""Report export: Markdown to .docx (headings, bullets, pipe tables, bold)."""

import io
import re

from docx import Document
from docx.shared import Pt

REPORT_KINDS = {
    "vm31_draft": "VM-31 style technical documentation draft (audience: appointed actuary, regulators)",
    "stakeholder_summary": "Plain-language results summary (audience: executives and non-actuarial stakeholders)",
}


def _add_runs(paragraph, text):
    for part in re.split(r"(\*\*[^*]+\*\*)", text):
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part:
            paragraph.add_run(part)


def markdown_to_docx(markdown, title):
    doc = Document()
    doc.styles["Normal"].font.size = Pt(10.5)
    doc.add_heading(title, level=0)
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                table = doc.add_table(rows=len(rows), cols=max(len(r) for r in rows))
                table.style = "Light Grid Accent 1"
                for r, row in enumerate(rows):
                    for c, cell in enumerate(row):
                        table.cell(r, c).text = ""
                        _add_runs(table.cell(r, c).paragraphs[0], cell)
            continue
        m = re.match(r"^(#{1,4})\s+(.*)", line)
        if m:
            doc.add_heading(m.group(2), level=min(len(m.group(1)), 4))
        elif re.match(r"^\s*[-*]\s+", line):
            _add_runs(doc.add_paragraph(style="List Bullet"), re.sub(r"^\s*[-*]\s+", "", line))
        elif re.match(r"^\s*\d+\.\s+", line):
            _add_runs(doc.add_paragraph(style="List Number"), re.sub(r"^\s*\d+\.\s+", "", line))
        elif line.strip():
            _add_runs(doc.add_paragraph(), line)
        i += 1
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
